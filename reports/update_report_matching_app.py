"""
Script cập nhật Báo Cáo Cuối Kỳ Đồ Án Tương Tác Dữ Liệu Trực Quan (Nhóm 22)
Khớp chuẩn xác 100% với Web Streamlit hiện tại (5 Mục Thuyết Trình, 8 Hero Charts, 50 Bang, 3D Globe, ROI Simulator, What-If Simulator)
Tác giả: Nhóm 22 - HCMUTE
"""

import os
import sys
try:
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
except Exception:
    pass
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

DOCX_IN = os.path.join(os.path.dirname(__file__), "..", "NHOM22_DOTRONGKHOI_BUIDUCHUY_TRUONGQUOCDUY_IDV_REPORT.docx")
DOCX_OUT_ROOT = os.path.join(os.path.dirname(__file__), "..", "NHOM22_DOTRONGKHOI_BUIDUCHUY_TRUONGQUOCDUY_IDV_REPORT.docx")
DOCX_OUT_REPORTS = os.path.join(os.path.dirname(__file__), "NHOM22_DOTRONGKHOI_BUIDUCHUY_TRUONGQUOCDUY_IDV_REPORT.docx")
PDF_OUT_ROOT = os.path.join(os.path.dirname(__file__), "..", "NHOM22_DOTRONGKHOI_BUIDUCHUY_TRUONGQUOCDUY_IDV_REPORT.pdf")
PDF_OUT_REPORTS = os.path.join(os.path.dirname(__file__), "NHOM22_DOTRONGKHOI_BUIDUCHUY_TRUONGQUOCDUY_IDV_REPORT.pdf")

def set_cell_text(cell, text, bold=False, italic=False, font_size_pt=10, font_name="Times New Roman", align=None, color_rgb=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    if color_rgb:
        run.font.color.rgb = color_rgb

def set_paragraph_text(p, text, bold=False, italic=False, font_size_pt=11, font_name="Times New Roman", line_spacing=1.15, space_before=Pt(3), space_after=Pt(4), align=None, color_rgb=None):
    p.text = ""
    if align:
        p.alignment = align
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = line_spacing
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    if color_rgb:
        run.font.color.rgb = color_rgb
    return p

def set_bullet_paragraph(p, bold_prefix, normal_text, font_size_pt=11, font_name="Times New Roman", space_before=Pt(3), space_after=Pt(3)):
    p.text = ""
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_bold = p.add_run(bold_prefix)
    r_bold.bold = True
    r_bold.font.name = font_name
    r_bold.font.size = Pt(font_size_pt)
    r_norm = p.add_run(normal_text)
    r_norm.font.name = font_name
    r_norm.font.size = Pt(font_size_pt)
    return p

def main():
    print(f"[*] Đang nạp tài liệu: {os.path.abspath(DOCX_IN)}...")
    doc = docx.Document(DOCX_IN)

    # 1. CẬP NHẬT TABLE 1: BÌA BÁO CÁO (DANH SÁCH THÀNH VIÊN)
    print("[1] Cập nhật Bảng thành viên trang bìa...")
    tbl_cover = doc.tables[0]
    members = [
        ("Trương Quốc Duy (Trưởng nhóm)", "24133009"),
        ("Đỗ Trọng Khôi", "20133056"),
        ("Bùi Đức Huy", "24133021")
    ]
    for idx, (name, mssv) in enumerate(members):
        set_cell_text(tbl_cover.cell(idx, 0), name, bold=True, font_size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(tbl_cover.cell(idx, 1), mssv, bold=True, font_size_pt=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    # 2. CẬP NHẬT TABLE 2: BẢNG PHÂN CÔNG NHIỆM VỤ THÀNH VIÊN
    print("[2] Cập nhật Bảng phân công nhiệm vụ (Table 2)...")
    tbl_roles = doc.tables[1]
    role_rows = [
        ("Trương Quốc Duy", "24133009", "Trưởng nhóm: Kiến trúc hệ thống tổng thể; Xây dựng Pipeline ETL dữ liệu & Làm sạch (nối 4 bảng, xử lý missing values, kiểm định ngoại lai IQR); Khám phá dữ liệu tĩnh EDA (10 biểu đồ); Huấn luyện mô hình Hồi quy Logistic, Feature Engineering & Đánh giá hiệu năng (ROC-AUC, Odds Ratio); Thiết kế & Lập trình toàn bộ Dashboard Streamlit (5 phân hệ, 8 Hero Charts, Bản đồ 3D Globe & US Flat Map, Drill-down 360°, What-If & ROI Simulator); Tối ưu hóa hiệu năng & Triển khai Cloud; Thuyết trình chính Kỹ thuật & Demo."),
        ("Đỗ Trọng Khôi", "20133056", "Thành viên: Soạn thảo, định dạng và tổng hợp toàn bộ Báo cáo tài liệu kỹ thuật Word (chuẩn IEEE / cấu trúc đồ án); Biên tập nội dung thuyết minh và đối chiếu Barem điểm; Tổng hợp tài liệu tham khảo và tài liệu hướng dẫn đồ án; Thuyết trình phần Cấu trúc báo cáo & Barem điểm."),
        ("Bùi Đức Huy", "24133021", "Thành viên: Khảo sát bối cảnh bài toán viễn thông, thu thập bộ dữ liệu Telco và mô tả Từ điển dữ liệu ban đầu (Data Dictionary); Chuẩn bị tài liệu & slide thuyết trình; Thuyết trình phần Mở đầu (Giới thiệu đề tài, mục tiêu nghiên cứu và tổng quan tập dữ liệu).")
    ]
    for idx, (name, mssv, role_desc) in enumerate(role_rows):
        r_idx = idx + 1
        set_cell_text(tbl_roles.cell(r_idx, 0), name, bold=True, font_size_pt=10)
        set_cell_text(tbl_roles.cell(r_idx, 1), mssv, bold=True, font_size_pt=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(tbl_roles.cell(r_idx, 2), role_desc, bold=False, font_size_pt=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # 3. CẬP NHẬT TABLE 3: BẢNG BAREM ĐÁNH GIÁ CHI TIẾT
    print("[3] Cập nhật Bảng Barem đánh giá (Table 3)...")
    tbl_barem = doc.tables[2]
    # Hàng 2.1 Thiết kế Dashboard (Row 5)
    set_cell_text(
        tbl_barem.cell(5, 3),
        "Streamlit + Plotly với 5 phân hệ tiến trình chuyên sâu, đúng 8 Hero Charts tinh hoa, Quả địa cầu 3D xoay 360° & US Flat Map 50 bang, 6 bộ lọc sidebar tương tác, và Drill-down hồ sơ 360 độ.",
        font_size_pt=9.5, align=WD_ALIGN_PARAGRAPH.LEFT
    )
    # Hàng 3.3 Trực quan Dự báo (Row 8)
    set_cell_text(
        tbl_barem.cell(8, 3),
        "Tích hợp biểu đồ Feature Importance (Hero Chart 6), đồng hồ Gauge xác suất Churn (Hero Chart 7) và công cụ What-If Simulator thời gian thực cùng dự phóng can thiệp trước/sau.",
        font_size_pt=9.5, align=WD_ALIGN_PARAGRAPH.LEFT
    )

    # 4. CẬP NHẬT CHƯƠNG 1.2 (Mục tiêu 3) & CHƯƠNG 1.5 (Tầng 5)
    print("[4] Cập nhật Chương 1.2 và 1.5...")
    for p in doc.paragraphs:
        if "Mục tiêu 3: Phát triển Bảng điều khiển tương tác" in p.text:
            set_paragraph_text(
                p,
                "Mục tiêu 3: Phát triển Bảng điều khiển tương tác (Interactive Dashboard): Sử dụng nền tảng Streamlit kết hợp với thư viện đồ họa động Plotly để tạo ra một không gian trực quan hóa hiện đại với 5 phân hệ chuyên sâu theo tiến trình thuyết trình, đúng 8 Hero Charts cốt lõi (bao gồm Quả địa cầu 3D xoay 360° & Bản đồ phẳng 50 bang, biểu đồ phân tích 2 chiều Hợp đồng x Mạng, đường cong duy trì thâm niên Retention Decay, ma trận 4 phân khúc Value vs Risk, đồng hồ Gauge What-If Simulator, biểu đồ kinh tế chiến dịch ROI Simulator). Cung cấp 6 bộ lọc sidebar đa chiều và tính năng khoan sâu (Drill-Down) truy xuất hồ sơ 360 độ của từng khách hàng. [6], [7]"
            )
        elif "Tầng 5 - Interactive Dashboard Presentation:" in p.text:
            set_paragraph_text(
                p,
                "Tầng 5 - Interactive Dashboard Presentation: Mã nguồn 'src/app.py' xây dựng giao diện Streamlit kết hợp Plotly Express/Graph Objects, cung cấp 5 phân hệ chuyên sâu theo tiến trình thuyết trình, đúng 8 Hero Charts cốt lõi, 6 bộ lọc dữ liệu đa chiều, drill-down hồ sơ khách hàng 360 độ và công cụ mô phỏng What-If Simulator cùng ROI Simulator."
            )

    # 5. CẬP NHẬT TABLE 21: PSEUDO-CODE LỌC ĐA CHIỀU (CHƯƠNG 4.1)
    print("[5] Cập nhật Pseudo-code lọc đa chiều (Table 21)...")
    tbl_code_dash = doc.tables[20]
    cell_code = tbl_code_dash.cell(0, 0)
    cell_code.text = ""
    p_code = cell_code.paragraphs[0]
    p_code.paragraph_format.space_before = Pt(3)
    p_code.paragraph_format.space_after = Pt(3)
    p_code.paragraph_format.line_spacing = 1.05
    code_text = """# Pseudo-code: Cơ chế Điều hướng 5 bước & Bộ lọc đa chiều trên Dashboard Streamlit
filtered_df = df_raw.copy()
if selected_cohort != "Toàn bộ chu kỳ (1 - 6 năm)":
    filtered_df = filter_by_cohort(filtered_df, selected_cohort)
if selected_state != "Tất cả 50 tiểu bang":
    filtered_df = filtered_df[filtered_df['State'] == selected_state]
if selected_internet != "Tất cả gói cước":
    filtered_df = filtered_df[filtered_df['InternetService'] == selected_internet]
if selected_contract != "Tất cả loại hợp đồng":
    filtered_df = filtered_df[filtered_df['Contract'] == selected_contract]
if selected_payment != "Tất cả phương thức":
    filtered_df = filtered_df[filtered_df['PaymentMethod'] == selected_payment]
filtered_df = filtered_df[(filtered_df['MonthlyCharges'] >= min_mrr) & (filtered_df['MonthlyCharges'] <= max_mrr)]

# Cập nhật tức thời 5 thẻ KPI toàn cục & Băng chuyền Tài chính C-Level
total_cust = len(filtered_df)
churn_rate = (filtered_df['Churn'] == 'Yes').mean() * 100
avg_arpu = filtered_df['MonthlyCharges'].mean()
total_clv = filtered_df['TotalCharges'].sum()
avg_csat = filtered_df['SatisfactionScore'].mean()
mrr_lost = filtered_df[filtered_df['Churn'] == 'Yes']['MonthlyCharges'].sum()
arr_lost = mrr_lost * 12
retained_pot_15 = arr_lost * 0.15"""
    r = p_code.add_run(code_text)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(226, 232, 240)

    # 6. CẬP NHẬT NGUYÊN LÝ SHNEIDERMAN (CHƯƠNG 4.2)
    print("[6] Cập nhật Nguyên lý Shneiderman (Chương 4.2)...")
    for p in doc.paragraphs:
        if "1. Overview first:" in p.text:
            set_paragraph_text(
                p,
                "1. Overview first: 5 thẻ chỉ số KPI toàn cục ở đầu trang (Tổng khách hàng: 7,043, Tỷ lệ Churn: 26.5%, ARPU: $64.76, Doanh thu tích lũy CLV: $16.06M, Điểm hài lòng CSAT: 3.53/5.0) kết hợp Băng chuyền Báo cáo Tài chính C-Level (Executive Financial Impact: Thất thoát MRR $139,131/tháng, Thất thoát ARR $1.67M/năm, Tiềm năng bảo vệ +$250,436/năm, Điểm nóng rủi ro 86.9% từ HĐ Tháng và 57.0% từ Cáp quang)."
            )
        elif "2. Zoom and filter:" in p.text:
            set_paragraph_text(
                p,
                "2. Zoom and filter: Thanh công cụ Sidebar tích hợp 6 bộ lọc tương tác đa chiều (Chu kỳ thâm niên theo năm, Địa bàn 50 bang viễn thông Hoa Kỳ, Gói dịch vụ Internet, Loại hợp đồng cam kết, Phương thức thanh toán, Thanh trượt khoảng cước phí hàng tháng USD) cùng nút đặt lại bộ lọc tức thời (Reset Filters)."
            )
        elif "3. Details-on-demand:" in p.text:
            set_paragraph_text(
                p,
                "3. Details-on-demand: Tab 5 cung cấp tính năng Drill-Down chuyên sâu: Chọn bất kỳ mã khách hàng nào (customerID), hệ thống lập tức mở rộng toàn bộ hồ sơ 360 độ gồm 4 khối thông tin (Nhân khẩu học, Hợp đồng & Thanh toán, Tài chính & Dịch vụ, Trải nghiệm & Phản hồi CSAT) kèm huy hiệu trạng thái rủi ro cá nhân hóa."
            )

    # 7. CẬP NHẬT TABLE 22: DANH MỤC BIỂU ĐỒ TRÊN DASHBOARD (CHƯƠNG 4.3)
    print("[7] Cập nhật Bảng danh mục biểu đồ Dashboard (Table 22)...")
    tbl_charts = doc.tables[21]
    dash_chart_data = [
        ("Tab 1: Tổng quan & Bản đồ Địa lý", "Thẻ KPI Metrics Toàn Cục", "Streamlit Metric Cards", "Tổng khách hàng (7,043), Churn Rate (26.5%), ARPU ($64.76), CLV ($16.06M), Điểm CSAT (3.53/5.0)."),
        ("Tab 1: Tổng quan & Bản đồ Địa lý", "Băng Chuyền Báo Cáo Tài Chính (Executive Financial Impact)", "Custom Financial Banner", "Định lượng: Thất thoát MRR ($139,131), ARR Loss ($1.67M), Tiềm năng bảo vệ (+$250,436), Điểm nóng rủi ro."),
        ("Tab 1: Tổng quan & Bản đồ Địa lý", "Hero Chart 1a: Tỷ Lệ Churn Tổng Thể & 1b: Phân Bố Theo Hợp Đồng", "Plotly Pie (hole=0.6) & Plotly Bar", "Tỷ trọng Ở lại (73.5%) vs Rời mạng (26.5%) và tỷ lệ rời mạng phân theo 3 kỳ hạn hợp đồng cam kết."),
        ("Tab 1: Tổng quan & Bản đồ Địa lý", "Hero Chart 2: Địa Cầu 3D Xoay 360° & Bản Đồ Phẳng 50 Bang", "Plotly Choropleth (3D Globe / 2D Map)", "Phân bố không gian 50 bang Hoa Kỳ (Animation xoay 360°, slider kinh độ -180° đến +180°, 5 preset góc nhìn, 4 chỉ số)."),
        ("Tab 1: Tổng quan & Bản đồ Địa lý", "Thẻ Tóm Tắt 3 Điểm Nóng Địa Lý", "Callout Badges", "Báo cáo nhanh: 🔴 Bang Churn cao nhất, 🟢 Bang an toàn nhất, và 🗺️ Độ lệch vùng (Spread)."),
        ("Tab 2: Chẩn đoán Dữ liệu & Hành vi", "Hero Chart 3: Tỷ Lệ Rời Mạng Theo Hợp Đồng & Gói Internet (%)", "Plotly Bar (Grouped Multi-dimension)", "Chẩn đoán 2 chiều: Khách hàng Gói tháng + Cáp quang Churn kỷ lục 54.6%, trong khi khách HĐ 2 năm < 8%."),
        ("Tab 2: Chẩn đoán Dữ liệu & Hành vi", "Hero Chart 4: Đường Cong Duy Trì Khách Hàng (Retention Decay)", "Plotly Line Chart (8 Cohorts)", "Theo dõi tỷ lệ gắn bó qua 8 mốc thâm niên (0-6th đến 61-72th), vạch rõ điểm gãy Retention Cliff sau năm đầu."),
        ("Tab 2: Chẩn đoán Dữ liệu & Hành vi", "Hero Chart 5: Ma Trận 4 Phân Khúc Giá Trị vs Rủi Ro", "Plotly Scatter (Ngưỡng VIP $70)", "Tương quan Cước tháng vs Thâm niên phân loại 4 nhóm: VIP Rủi Ro Cao, VIP Trung Thành, Phổ Thông Rủi Ro, Phổ Thông Ổn Định."),
        ("Tab 2: Chẩn đoán Dữ liệu & Hành vi", "Thẻ Insights Chẩn Đoán Cốt Lõi", "Callout Text Cards", "Đưa ra phát hiện nguyên nhân gốc rễ (Root Cause) và chỉ dẫn can thiệp ưu tiên nhóm VIP Rủi Ro Cao."),
        ("Tab 3: Dự báo AI & Simulator", "Hero Chart 6: Trọng Số Các Yếu Tố Quyết Định Churn", "Plotly Bar (Horizontal)", "Trọng số hệ số hồi quy Log-Odds (β) và tỷ số chênh Odds Ratio phân biệt yếu tố tăng nguy cơ vs giữ chân."),
        ("Tab 3: Dự báo AI & Simulator", "Thẻ Hiệu Năng Mô Hình Trên Tập Kiểm Định (Test Set)", "Streamlit Metric Cards", "Thông số kiểm định mô hình Logistic Regression: Accuracy 80.77%, ROC-AUC 0.8421, F1-Score 60.89%."),
        ("Tab 3: Dự báo AI & Simulator", "Trình Mô Phỏng Nguy Cơ Rời Mạng (What-If Real-Time Simulator)", "Streamlit Form + Predict Engine", "Nhập hồ sơ hợp đồng linh hoạt 8 trường để mô hình máy học tính toán tức thời xác suất Churn."),
        ("Tab 3: Dự báo AI & Simulator", "Hero Chart 7: Đồng Hồ Đo Xác Suất Rời Mạng (Gauge Chart)", "Plotly Gauge Indicator & Metric Card", "Đồng hồ kim Gauge đo xác suất Churn theo 3 dải màu rủi ro (<35%, 35-60%, >=60%) và lượng hóa cước năm bị ảnh hưởng."),
        ("Tab 3: Dự báo AI & Simulator", "Card Dự Đoán Tác Động Sau Can Thiệp (Trước & Sau Giữ Chân)", "Metric Container Cards", "Lập tức lượng hóa mức giảm xác suất Churn khi đổi HĐ 1 năm, thêm TechSupport 24/7, hoặc Combo Toàn Diện."),
        ("Tab 4: Khuyến nghị Chiến lược & ROI", "Khung Chiến Lược Hành Động 3 Trụ Cột (Actionable Playbook)", "Custom Strategy Cards Layout", "Đề xuất giải pháp cụ thể: Trụ cột 1 (Fiber Shield), Trụ cột 2 (Chuyển sang Auto-Pay), Trụ cột 3 (Khóa hợp đồng)."),
        ("Tab 4: Khuyến nghị Chiến lược & ROI", "Trình Mô Phỏng Chiến Dịch Giữ Chân (ROI Simulator)", "Streamlit Sliders + Dynamic Math", "Thiết lập tham số (Tỷ lệ tiếp cận, Tỷ lệ giữ chân, Ngân sách) dự phóng kết quả tài chính C-Level thời gian thực."),
        ("Tab 4: Khuyến nghị Chiến lược & ROI", "Hero Chart 8: So Sánh Hiệu Quả Kinh Tế Chiến Dịch (Quy Năm USD)", "Plotly Bar / Column Chart", "Trực quan hóa 4 định mức tài chính: Tổn thất ban đầu, Ngân sách đầu tư, Doanh thu bảo vệ được và Lợi nhuận ròng."),
        ("Tab 4: Khuyến nghị Chiến lược & ROI", "Báo Cáo Tóm Tắt Chiến Lược Điều Hành (Executive Brief)", "Streamlit Download Button", "Nút tải trực tiếp file báo cáo tóm tắt chiến lược dành cho Ban Điều Hành (.MD)."),
        ("Tab 5: Kiến trúc Dữ liệu & Tra cứu 360°", "Thước Đo Chất Lượng Dữ Liệu (Data Quality SLA)", "Streamlit Metric Cards", "Cam kết độ tin cậy: Tỷ lệ khớp nối 100.0%, Missing Values 0.00%, Độ trễ truy vấn < 15ms, 39 Đặc trưng."),
        ("Tab 5: Kiến trúc Dữ liệu & Tra cứu 360°", "Tra Cứu Drill-Down Hồ Sơ 360° Khách Hàng Cụ Thể", "Selectbox + Profile Card 360", "Chọn CustomerID hiển thị toàn diện: Nhân khẩu học, Hợp đồng & Thanh toán, Tài chính, CSAT & Lý do rời mạng."),
        ("Tab 5: Kiến trúc Dữ liệu & Tra cứu 360°", "Bảng Chi Tiết Khách Hàng & Nút Xuất File CSV", "Interactive Table & Download Button", "Bảng dữ liệu tương tác đầy đủ định dạng trực quan và xuất báo cáo CSV tập khách hàng đã lọc.")
    ]
    for idx, row in enumerate(dash_chart_data):
        r_idx = idx + 1
        set_cell_text(tbl_charts.cell(r_idx, 0), row[0], bold=True, font_size_pt=9.5)
        set_cell_text(tbl_charts.cell(r_idx, 1), row[1], bold=False, font_size_pt=9.5)
        set_cell_text(tbl_charts.cell(r_idx, 2), row[2], bold=False, font_size_pt=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(tbl_charts.cell(r_idx, 3), row[3], bold=False, font_size_pt=9, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # 8. CẬP NHẬT CAPTION CỦA 6 HÌNH ẢNH SCREENSHOTS DASHBOARD (HÌNH 11 ĐẾN 16)
    print("[8] Cập nhật Caption các hình ảnh Dashboard...")
    caption_map = {
        "Hình 11:": "Hình 11: Giao diện của Tab 1: Tổng quan và Bản đồ địa lý (Hero Charts 1 & 2)",
        "Hình 12:": "Hình 12: Giao diện của Tab 2: Chẩn đoán dữ liệu và hành vi (Hero Charts 3, 4 & 5)",
        "Hình 13:": "Hình 13: Giao diện của Tab 3: Dự đoán AI và mô phỏng – Biểu đồ trọng số các yếu tố quyết định Churn (Hero Chart 6)",
        "Hình 14:": "Hình 14: Giao diện của Tab 3: Dự đoán AI và mô phỏng – Trình mô phỏng What-If Simulator & Đồng hồ Gauge (Hero Chart 7)",
        "Hình 15:": "Hình 15: Giao diện của Tab 4: Khuyến nghị chiến lược Prescriptive & Trình mô phỏng ROI (Hero Chart 8)",
        "Hình 16:": "Hình 16: Giao diện của Tab 5: Kiến trúc dữ liệu SLA & Tra cứu hồ sơ khách hàng 360°"
    }
    for p in doc.paragraphs:
        for prefix, new_cap in caption_map.items():
            if prefix in p.text and "Giao diện của Tab" in p.text:
                set_paragraph_text(p, new_cap, italic=True, font_size_pt=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color_rgb=RGBColor(71, 85, 105))

    # 9. CẬP NHẬT CHƯƠNG 4.4: BẢN ĐỒ ĐỊA LÝ KHÔNG GIAN
    print("[9] Cập nhật Chương 4.4: Bản đồ địa lý không gian...")
    for idx, p in enumerate(doc.paragraphs):
        if "4.4 Phân Tích Không Gian Địa Lý" in p.text:
            p_next = doc.paragraphs[idx+1]
            set_paragraph_text(
                p_next,
                "Bản đồ địa lý không gian (Hero Chart 2) tại Tab 1: Tổng quan & Bản đồ Địa lý trực quan hóa toàn diện phân bổ khách hàng và nguy cơ Churn trên toàn bộ 50 tiểu bang Hoa Kỳ, hỗ trợ chuyển đổi linh hoạt giữa hai chế độ trực quan hóa đỉnh cao:"
            )
        elif "Kích thước bong bóng (Size):" in p.text:
            set_paragraph_text(
                p,
                "Quả địa cầu 3D xoay 360° (3D Interactive Globe): Sử dụng phép chiếu trực giao orthographic với hiệu ứng bề mặt đại dương và lục địa sống động. Tích hợp animation xoay quanh trục tự động (nút Play/Pause), thanh trượt tinh chỉnh góc kinh độ (-180° đến +180°) và 5 nút preset góc nhìn nhanh (🇺🇸 Toàn cảnh Hoa Kỳ -98°, 🗽 Bờ Đông -75°, 🌉 Bờ Tây -125°, 🇻🇳 Châu Á +105°, 🌍 Châu Âu 0°)."
            )
        elif "Thang màu sắc (Color Scale):" in p.text:
            set_paragraph_text(
                p,
                "Bản đồ phẳng 2D Choropleth: Thể hiện trực quan ranh giới 50 tiểu bang Hoa Kỳ với dải màu sắc thái chuyển đổi liên tục theo độ lớn của từng chỉ số."
            )
        elif "Dữ liệu tương tác khi di chuột (Hover Data):" in p.text:
            set_paragraph_text(
                p,
                "Lựa chọn 4 chỉ số nghiệp vụ linh hoạt & Thẻ tóm tắt 3 điểm nóng: Cho phép nhà quản lý phân tích không gian theo 4 lăng kính: Tỷ lệ Churn (%), Quy mô Khách hàng, Cước phí TB ($), và Tổng Doanh thu CLV ($). Đồng thời báo cáo tức thời về: 🔴 Tiểu bang có tỷ lệ Churn cao nhất, 🟢 Tiểu bang an toàn nhất, và 🗺️ Độ lệch vùng (Spread)."
            )

    # 10. CẬP NHẬT CHƯƠNG 4.5: KIẾN TRÚC DỮ LIỆU SLA & TRA CỨU 360°
    print("[10] Cập nhật Chương 4.5: Kiến trúc dữ liệu SLA & Tra cứu 360°...")
    for idx, p in enumerate(doc.paragraphs):
        if "4.5 Tính Năng Drill-Down Hồ Sơ 360 Độ" in p.text:
            set_paragraph_text(p, "4.5 Kiến Trúc Dữ Liệu SLA & Tính Năng Drill-Down Hồ Sơ 360 Độ (Tab 5)", bold=True, font_size_pt=13, color_rgb=RGBColor(30, 58, 138))
            p_next = doc.paragraphs[idx+1]
            set_paragraph_text(
                p_next,
                "Tại Tab 5, hệ thống minh chứng năng lực kỹ thuật dữ liệu cấp doanh nghiệp thông qua sơ đồ Enterprise Data Pipeline 5 giai đoạn và bảng cam kết chất lượng dịch vụ SLA: Tỷ lệ khớp nối quan hệ 100.0% (7,043 bản ghi), Dữ liệu rỗng 0.00% (xử lý triệt để 11 giá trị TotalCharges), Độ trễ truy vấn < 15 ms qua bộ đệm nhớ Streamlit Cache, và 39 đặc trưng chuẩn hóa. Đồng thời, tính năng Drill-Down cho phép người dùng chọn bất kỳ CustomerID nào để lập tức hiển thị thẻ hồ sơ 360 độ gồm 4 phân khu (Nhân khẩu học, Hợp đồng & Thanh toán, Tài chính & Dịch vụ, Trải nghiệm CSAT & Lý do rời mạng) kèm nút xuất dữ liệu đã lọc sang file CSV."
            )

    # 11. CẬP NHẬT CHƯƠNG 4.8: KHUYẾN NGHỊ CHIẾN LƯỢC & ROI SIMULATOR (TAB 4)
    print("[11] Cập nhật Chương 4.8: Chiến lược & ROI Simulator...")
    for idx, p in enumerate(doc.paragraphs):
        if "1. Contract Migration:" in p.text:
            set_paragraph_text(
                p,
                "Trụ cột 1 - Sản phẩm & Mạng lưới (Gói 'Fiber Shield Bundle'): Đóng gói miễn phí 3-6 tháng dịch vụ an ninh mạng OnlineSecurity và hỗ trợ kỹ thuật TechSupport 24/7 vào gói Cáp quang. Dữ liệu thực nghiệm chứng minh tỷ lệ Churn của khách có 2 gói này giảm ngoạn mục xuống còn 15.8% (so với 54.6% ở nhóm gói tháng cáp quang đơn lẻ)."
            )
        elif "2. VAS Bundling:" in p.text:
            set_paragraph_text(
                p,
                "Trụ cột 2 - Tài chính & Thanh toán (Chuyển đổi sang Auto-Pay): Khắc phục ma sát thanh toán bằng chính sách tặng voucher giảm $5/tháng trong 3 tháng liên tiếp cho khách hàng chuyển từ Séc điện tử (Electronic check) sang Bank Transfer hoặc Credit Card tự động, giúp giảm tỷ lệ Churn từ 45.3% xuống chỉ còn 16.7%."
            )
        elif "3. Auto-Pay Incentive:" in p.text:
            set_paragraph_text(
                p,
                "Trụ cột 3 - Vòng đời & Hợp đồng (Khóa hợp đồng & Vượt vùng tử thần): Chủ động tiếp cận khách hàng vào tháng thứ 4 & tháng thứ 8 để xử lý khiếu nại CSAT, kèm chính sách chiết khấu 12% cước năm khi cam kết chuyển sang hợp đồng 1-2 năm, triệt tiêu 65% ca Churn tập trung trong năm đầu."
            )
        elif "4. Early Warning CSKH:" in p.text:
            set_paragraph_text(
                p,
                "Trình mô phỏng ROI Simulator & Hero Chart 8: Công cụ giả lập kinh tế dành cho Ban Điều Hành (C-Level) với 3 thanh trượt tham số (Tỷ lệ tiếp cận CSKH, Tỷ lệ giữ chân thành công, Ngân sách ưu đãi/khách) lập tức tính toán 4 chỉ số tài chính: Số khách giữ chân thành công, Doanh thu bảo vệ được (ARR Preserved), Ngân sách chiến dịch, và Lợi nhuận ròng mang lại. Hero Chart 8 trực quan hóa trực tiếp 4 cột tài chính tương phản."
            )
        elif "Ước tính hiệu quả tài chính (Financial ROI):" in p.text:
            set_paragraph_text(
                p,
                "Ước tính hiệu quả tài chính (Financial ROI) & Báo cáo điều hành (.MD): Với kịch bản tiếp cận 40% tệp rủi ro và giữ chân thành công 25%, nhà mạng giữ lại được 186 khách hàng, bảo toàn $166,153 USD doanh thu quy năm với ngân sách đầu tư chỉ $8,964 USD, mang lại lợi nhuận ròng $157,189 USD và ROI vượt 1,000%. Nút tải Executive Brief (.MD) cho phép xuất ngay báo cáo tóm tắt dành cho lãnh đạo."
            )

    # 12. CẬP NHẬT CHƯƠNG 5.5: TRÌNH MÔ PHỎNG WHAT-IF SIMULATOR (TAB 3)
    print("[12] Cập nhật Chương 5.5: What-If Simulator (Tab 3)...")
    for idx, p in enumerate(doc.paragraphs):
        if "5.5 Tích Hợp Mô Hình Dự Báo Trực Quan" in p.text:
            p_next = doc.paragraphs[idx+1]
            set_paragraph_text(
                p_next,
                "Điểm nổi bật của đề tài là việc tích hợp trực tiếp Pipeline máy học Scikit-Learn vào Dashboard Streamlit tại Tab 3: Dự báo AI & Simulator. Người dùng có thể điều chỉnh 8 thông số hợp đồng giả định (Thâm niên, Cước phí, Hợp đồng, Gói Internet, Phương thức thanh toán, OnlineSecurity, TechSupport, PaperlessBilling) để mô hình tính toán tức thời xác suất Churn. Kết quả được biểu diễn sống động qua Hero Chart 7 (Đồng hồ Gauge với 3 dải màu rủi ro và ngưỡng 60%), thẻ ước tính cước năm bị ảnh hưởng ($/năm) và khối Card dự đoán tác động giảm nguy cơ rời mạng khi áp dụng 3 phương án giữ chân (Đổi HĐ 1 năm, Tặng TechSupport 24/7, Gói Combo Toàn Diện)."
            )

    # 13. CẬP NHẬT CHƯƠNG 6.1: ĐÁNH GIÁ KẾT QUẢ ĐẠT ĐƯỢC
    print("[13] Cập nhật Chương 6.1: Kết luận...")
    for idx, p in enumerate(doc.paragraphs):
        if "Về mặt trực quan hóa và tương tác, nhóm đã xây dựng 10 biểu đồ tĩnh" in p.text:
            set_paragraph_text(
                p,
                "Về mặt trực quan hóa và tương tác, nhóm đã xây dựng 10 biểu đồ tĩnh phục vụ phân tích khám phá dữ liệu (EDA), đồng thời phát triển thành công Dashboard tương tác trên nền tảng Streamlit tích hợp bộ thư viện Plotly với 5 phân hệ tiến trình chuyên sâu, đúng 8 Hero Charts cốt lõi cùng Quả địa cầu 3D xoay 360° và Bản đồ phẳng 50 bang Hoa Kỳ. Cuối cùng, mô hình Hồi quy Logistic (Logistic Regression) được triển khai thử nghiệm đạt chỉ số hiệu năng ấn tượng với ROC-AUC đạt 0.8421, khẳng định khả năng dự đoán chính xác tỷ lệ rời bỏ của khách hàng."
            )

    # LƯU FILE DOCX
    print(f"[*] Đang lưu file Word cập nhật vào '{DOCX_OUT_ROOT}'...")
    doc.save(DOCX_OUT_ROOT)
    doc.save(DOCX_OUT_REPORTS)
    print("[+] ĐÃ LƯU THÀNH CÔNG CẢ 2 ĐƯỜNG DẪN DOCX!")

    # XUẤT FILE PDF TỰ ĐỘNG QUA WORD COM
    print("[*] Đang xuất file PDF qua Microsoft Word...")
    try:
        import win32com.client
        import shutil
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc_com = word.Documents.Open(os.path.abspath(DOCX_OUT_ROOT))
        doc_com.SaveAs(os.path.abspath(PDF_OUT_ROOT), FileFormat=17) # 17 = wdFormatPDF
        doc_com.Close()
        word.Quit()
        shutil.copyfile(PDF_OUT_ROOT, PDF_OUT_REPORTS)
        print("[+] ĐÃ XUẤT THÀNH CÔNG CẢ 2 ĐƯỜNG DẪN PDF!")
    except Exception as e:
        print(f"[!] Cảnh báo xuất PDF: {e}")

if __name__ == "__main__":
    main()
